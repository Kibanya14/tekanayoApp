# ============================================================================
# CATALOGUE IMPORT MANAGER - Importer des produits depuis fichiers
# ============================================================================
# Formats supportés: CSV, XLSX, DOCX, PDF, TXT, TSV
# Colonnes: name, price, category/category_id, description, quantity, etc.
# ============================================================================

import os
import io
import csv
import json
from typing import List, Dict, Tuple, Optional
from datetime import datetime

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None


class CatalogImportManager:
    """Gestionnaire pour importer un catalogue de produits"""
    
    # Colonnes reconnues
    RECOGNIZED_COLUMNS = {
        'name', 'product_name', 'titre',
        'price', 'prix', 'product_price',
        'category', 'categorie', 'category_id',
        'description', 'descriptif',
        'quantity', 'stock', 'quantite',
        'compare_price', 'prix_compare',
        'is_active', 'actif', 'enabled',
        'is_featured', 'en_vedette', 'featured'
    }
    
    # Mappage des colonnes variantes
    COLUMN_MAPPING = {
        'product_name': 'name',
        'titre': 'name',
        'product_price': 'price',
        'prix': 'price',
        'categorie': 'category',
        'category_id': 'category',
        'descriptif': 'description',
        'stock': 'quantity',
        'quantite': 'quantity',
        'prix_compare': 'compare_price',
        'actif': 'is_active',
        'enabled': 'is_active',
        'en_vedette': 'is_featured',
        'featured': 'is_featured',
    }
    
    @staticmethod
    def normalize_column_name(col_name: str) -> Optional[str]:
        """Normalise un nom de colonne"""
        if not col_name:
            return None
        col = col_name.strip().lower().replace(' ', '_')
        if col in CatalogImportManager.RECOGNIZED_COLUMNS:
            return col
        if col in CatalogImportManager.COLUMN_MAPPING:
            return CatalogImportManager.COLUMN_MAPPING[col]
        return None
    
    @staticmethod
    def parse_csv(file_content: bytes, encoding: str = 'utf-8') -> Tuple[List[str], List[Dict]]:
        """Parse un fichier CSV"""
        try:
            text = file_content.decode(encoding)
        except UnicodeDecodeError:
            try:
                text = file_content.decode('latin-1')
            except:
                return [], []
        
        lines = text.strip().split('\n')
        if not lines:
            return [], []
        
        reader = csv.DictReader(io.StringIO(text))
        rows = []
        
        for row in reader:
            if row:
                rows.append(row)
        
        return list(reader.fieldnames or []), rows
    
    @staticmethod
    def parse_xlsx(file_content: bytes) -> Tuple[List[str], List[Dict]]:
        """Parse un fichier XLSX (Excel)"""
        if not openpyxl:
            return [], []
        
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(file_content))
            worksheet = workbook.active
            
            # Récupérer les en-têtes
            headers = []
            for cell in worksheet[1]:
                headers.append(cell.value or '')
            
            if not headers:
                return [], []
            
            # Récupérer les données
            rows = []
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                row_dict = {}
                for i, header in enumerate(headers):
                    if i < len(row):
                        row_dict[header] = row[i]
                rows.append(row_dict)
            
            return headers, rows
        except Exception as e:
            print(f"Erreur XLSX: {e}")
            return [], []
    
    @staticmethod
    def parse_docx(file_content: bytes) -> Tuple[List[str], List[Dict]]:
        """Parse un fichier DOCX (Word)"""
        if not DocxDocument:
            return [], []
        
        try:
            doc = DocxDocument(io.BytesIO(file_content))
            rows = []
            
            # Chercher une table dans le document
            for table in doc.tables:
                headers = []
                for cell in table.rows[0].cells:
                    headers.append(cell.text.strip())
                
                if not headers:
                    continue
                
                for row in table.rows[1:]:
                    row_dict = {}
                    for i, header in enumerate(headers):
                        if i < len(row.cells):
                            row_dict[header] = row.cells[i].text.strip()
                    rows.append(row_dict)
                
                return headers, rows
            
            return [], []
        except Exception as e:
            print(f"Erreur DOCX: {e}")
            return [], []
    
    @staticmethod
    def parse_pdf(file_content: bytes) -> Tuple[List[str], List[Dict]]:
        """Parse un fichier PDF"""
        if not PyPDF2:
            return [], []
        
        try:
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text()
            
            # Parser comme CSV simple (heuristique)
            lines = text.strip().split('\n')
            rows = []
            
            if lines:
                # Première ligne = en-têtes
                headers = [h.strip() for h in lines[0].split('\t') or lines[0].split(',')]
                
                for line in lines[1:]:
                    parts = line.split('\t') or line.split(',')
                    row_dict = {}
                    for i, header in enumerate(headers):
                        if i < len(parts):
                            row_dict[header] = parts[i].strip()
                    rows.append(row_dict)
            
            return headers if lines else [], rows
        except Exception as e:
            print(f"Erreur PDF: {e}")
            return [], []
    
    @staticmethod
    def parse_tsv(file_content: bytes) -> Tuple[List[str], List[Dict]]:
        """Parse un fichier TSV"""
        try:
            text = file_content.decode('utf-8')
        except:
            text = file_content.decode('latin-1')
        
        lines = text.strip().split('\n')
        if not lines:
            return [], []
        
        # Première ligne = en-têtes
        headers = lines[0].split('\t')
        rows = []
        
        for line in lines[1:]:
            parts = line.split('\t')
            row_dict = {}
            for i, header in enumerate(headers):
                if i < len(parts):
                    row_dict[header] = parts[i]
            rows.append(row_dict)
        
        return headers, rows
    
    @staticmethod
    def parse_txt(file_content: bytes) -> Tuple[List[str], List[Dict]]:
        """Parse un fichier TXT"""
        try:
            text = file_content.decode('utf-8')
        except:
            text = file_content.decode('latin-1')
        
        lines = text.strip().split('\n')
        if not lines:
            return [], []
        
        # Essayer CSV avec délimiteur auto-détecté
        if ',' in lines[0]:
            return CatalogImportManager.parse_csv(file_content)
        elif '\t' in lines[0]:
            return CatalogImportManager.parse_tsv(file_content)
        else:
            # Ligne unique par produit
            return ['name', 'quantity'], [{'name': line.strip(), 'quantity': 1} for line in lines]
    
    @staticmethod
    def parse_file(file_obj) -> Tuple[List[str], List[Dict]]:
        """Parse un fichier uploadé"""
        filename = file_obj.filename.lower()
        file_content = file_obj.read()
        
        if filename.endswith('.csv'):
            return CatalogImportManager.parse_csv(file_content)
        elif filename.endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
            return CatalogImportManager.parse_xlsx(file_content)
        elif filename.endswith('.docx'):
            return CatalogImportManager.parse_docx(file_content)
        elif filename.endswith('.pdf'):
            return CatalogImportManager.parse_pdf(file_content)
        elif filename.endswith('.tsv'):
            return CatalogImportManager.parse_tsv(file_content)
        elif filename.endswith('.txt'):
            return CatalogImportManager.parse_txt(file_content)
        
        return [], []
    
    @staticmethod
    def normalize_data(headers: List[str], rows: List[Dict]) -> Tuple[List[str], List[Dict]]:
        """Normalise les en-têtes et données"""
        normalized_headers = []
        header_mapping = {}
        
        for header in headers:
            normalized = CatalogImportManager.normalize_column_name(header)
            if normalized:
                normalized_headers.append(normalized)
                header_mapping[header] = normalized
        
        # Normaliser les lignes
        normalized_rows = []
        for row in rows:
            normalized_row = {}
            for original_col, value in row.items():
                if original_col in header_mapping:
                    normalized_col = header_mapping[original_col]
                    normalized_row[normalized_col] = value
            
            if normalized_row:  # Si la ligne a au moins une colonne
                normalized_rows.append(normalized_row)
        
        return normalized_headers, normalized_rows
    
    @staticmethod
    def validate_and_convert_row(row: Dict) -> Tuple[bool, str, Dict]:
        """Valide et convertit une ligne de produit"""
        errors = []
        
        # Vérifier les champs requis
        name = (row.get('name') or '').strip()
        if not name:
            errors.append("Nom du produit manquant")
        
        # Prix
        price_str = row.get('price', '')
        try:
            price = float(price_str) if price_str else 0.0
            if price < 0:
                errors.append(f"Prix invalide: {price_str}")
                price = 0.0
        except (ValueError, TypeError):
            errors.append(f"Prix invalide: {price_str}")
            price = 0.0
        
        # Catégorie
        category = (row.get('category') or 'Autres').strip()
        if not category:
            category = 'Autres'
        
        # Description
        description = (row.get('description') or '').strip()
        
        # Quantité
        quantity_str = row.get('quantity', '')
        try:
            quantity = int(float(quantity_str)) if quantity_str else 0
            if quantity < 0:
                quantity = 0
        except (ValueError, TypeError):
            quantity = 0
        
        # Compare price
        compare_price_str = row.get('compare_price', '')
        try:
            compare_price = float(compare_price_str) if compare_price_str else None
        except (ValueError, TypeError):
            compare_price = None
        
        # Flags
        is_active = str(row.get('is_active', 'true')).lower() in ('true', '1', 'oui', 'yes', 'yes', 'actif')
        is_featured = str(row.get('is_featured', 'false')).lower() in ('true', '1', 'oui', 'yes', 'featured')
        
        if errors:
            return False, " | ".join(errors), {}
        
        return True, "", {
            'name': name,
            'price': price,
            'category': category,
            'description': description,
            'quantity': quantity,
            'compare_price': compare_price,
            'is_active': is_active,
            'is_featured': is_featured,
        }


# ============================================================================
# FONCTION DE ROUTE - À AJOUTER DANS backend/apps.py
# ============================================================================

def setup_catalog_import_routes(app):
    """Configure les routes d'import de catalogue"""
    from flask import request, flash, redirect, url_for, jsonify
    from backend.models import db, SellerProduct, SellerShop, SellerAdmin
    from flask import session, g
    
    @app.route("/admin/catalog/import", methods=["POST"])
    def admin_import_catalog():
        """Import catalogue admin - Pour boutique portail"""
        if not request.files.get('file'):
            flash("Aucun fichier sélectionné", "error")
            return redirect(url_for("tekanayo_admin_products"))
        
        file = request.files['file']
        if not file.filename:
            flash("Fichier invalide", "error")
            return redirect(url_for("tekanayo_admin_products"))
        
        # Parser le fichier
        headers, rows = CatalogImportManager.parse_file(file)
        if not rows:
            flash("Fichier vide ou format non supporté", "error")
            return redirect(url_for("tekanayo_admin_products"))
        
        # Normaliser
        headers, rows = CatalogImportManager.normalize_data(headers, rows)
        
        # Importer
        success_count = 0
        error_count = 0
        errors = []
        
        # Boutique portail
        shop = SellerShop.query.filter_by(is_portal_shop=True).first()
        if not shop:
            flash("Boutique portail introuvable", "error")
            return redirect(url_for("tekanayo_admin_products"))
        
        for i, row in enumerate(rows, 1):
            is_valid, error_msg, data = CatalogImportManager.validate_and_convert_row(row)
            
            if not is_valid:
                error_count += 1
                errors.append(f"Ligne {i}: {error_msg}")
                continue
            
            try:
                product = SellerProduct(
                    shop_id=shop.id,
                    name=data['name'],
                    price=data['price'],
                    category=data['category'],
                    description=data['description'],
                    quantity=data['quantity'],
                    compare_price=data['compare_price'],
                    is_active=data['is_active'],
                    is_featured=data['is_featured'],
                )
                db.session.add(product)
                success_count += 1
            except Exception as e:
                error_count += 1
                errors.append(f"Ligne {i}: Erreur BD - {str(e)}")
        
        db.session.commit()
        
        flash(f"✅ Import complété: {success_count} produit(s) ajouté(s)", "success")
        if errors:
            flash(f"⚠️ {error_count} erreur(s):\n" + "\n".join(errors[:5]), "warning")
        
        return redirect(url_for("tekanayo_admin_products"))
    
    @app.route("/vendeur/<slug>/catalog/import", methods=["POST"])
    def seller_import_catalog(slug):
        """Import catalogue vendeur"""
        from decorators import seller_session_required
        
        shop = SellerShop.query.filter_by(slug=slug).first_or_404()
        
        if not request.files.get('file'):
            flash("Aucun fichier sélectionné", "error")
            return redirect(url_for("seller_products_page", slug=slug))
        
        file = request.files['file']
        if not file.filename:
            flash("Fichier invalide", "error")
            return redirect(url_for("seller_products_page", slug=slug))
        
        # Vérifier permission
        current_admin = SellerAdmin.query.filter_by(
            id=session.get("seller_admin_id"),
            shop_id=shop.id
        ).first_or_404()
        
        if not (current_admin.is_owner or 'manage_products' in (current_admin.permissions or '').split(',')):
            flash("Permission refusée", "error")
            return redirect(url_for("seller_products_page", slug=slug))
        
        # Parser le fichier
        headers, rows = CatalogImportManager.parse_file(file)
        if not rows:
            flash("Fichier vide ou format non supporté", "error")
            return redirect(url_for("seller_products_page", slug=slug))
        
        # Normaliser
        headers, rows = CatalogImportManager.normalize_data(headers, rows)
        
        # Importer
        success_count = 0
        error_count = 0
        errors = []
        
        for i, row in enumerate(rows, 1):
            is_valid, error_msg, data = CatalogImportManager.validate_and_convert_row(row)
            
            if not is_valid:
                error_count += 1
                errors.append(f"Ligne {i}: {error_msg}")
                continue
            
            try:
                product = SellerProduct(
                    shop_id=shop.id,
                    name=data['name'],
                    price=data['price'],
                    category=data['category'],
                    description=data['description'],
                    quantity=data['quantity'],
                    compare_price=data['compare_price'],
                    is_active=data['is_active'],
                    is_featured=data['is_featured'],
                )
                db.session.add(product)
                success_count += 1
            except Exception as e:
                error_count += 1
                errors.append(f"Ligne {i}: Erreur BD - {str(e)}")
        
        db.session.commit()
        
        flash(f"✅ Import complété: {success_count} produit(s) ajouté(s)", "success")
        if errors:
            flash(f"⚠️ {error_count} erreur(s):\n" + "\n".join(errors[:5]), "warning")
        
        return redirect(url_for("seller_products_page", slug=slug))
